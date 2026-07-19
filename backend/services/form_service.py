"""
GovBrief AI — Form Generation Service
Generates filled DOCX/PDF forms from OCR-extracted data
"""
import os
from pathlib import Path
from datetime import datetime
from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from config import settings


FORM_TEMPLATES = {
    "pm_kisan": "PM-KISAN Application Form",
    "solar_subsidy": "PM Kusum Solar Pump Subsidy",
    "crop_insurance": "PMFBY Crop Insurance Form",
    "kisan_credit_card": "Kisan Credit Card Application",
    "soil_health": "Soil Health Card Request",
}


def generate_form_docx(form_type: str, data: dict) -> str:
    """Generate a filled DOCX form and return the filename."""
    doc = DocxDocument()

    # Header
    header = doc.add_heading(FORM_TEMPLATES.get(form_type, "Government Application Form"), 0)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph("Ministry of Agriculture & Farmers Welfare, Government of India")
    doc.add_paragraph("_" * 60)

    # Applicant Details Section
    doc.add_heading("Section A: Applicant Details", level=1)
    fields = [
        ("Full Name", data.get("name", "")),
        ("Father's Name", data.get("father_name", "")),
        ("Date of Birth", data.get("dob", "")),
        ("Aadhaar Number", data.get("id_number", "")),
        ("Mobile Number", data.get("phone", "")),
        ("State", data.get("state", "")),
        ("District", data.get("district", "")),
        ("Village / Town", data.get("village", "")),
        ("PIN Code", data.get("pincode", "")),
        ("Full Address", data.get("address", "")),
    ]
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(fields):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = str(value)

    doc.add_paragraph("")

    # Land Details
    doc.add_heading("Section B: Land Details", level=1)
    land_fields = [
        ("Survey Number / Khasra", data.get("survey_number", "")),
        ("Land Area (in Acres)", data.get("land_area", "")),
        ("Land Type", data.get("land_type", "Agricultural")),
        ("Crop Grown", data.get("crop_types", "")),
    ]
    table2 = doc.add_table(rows=len(land_fields), cols=2)
    table2.style = "Table Grid"
    for i, (label, value) in enumerate(land_fields):
        table2.rows[i].cells[0].text = label
        table2.rows[i].cells[1].text = str(value)

    doc.add_paragraph("")

    # Bank Details
    doc.add_heading("Section C: Bank Details", level=1)
    bank_fields = [
        ("Bank Name", data.get("bank_name", "")),
        ("Account Number", data.get("bank_account", "")),
        ("IFSC Code", data.get("ifsc_code", "")),
        ("Branch Name", data.get("branch_name", "")),
    ]
    table3 = doc.add_table(rows=len(bank_fields), cols=2)
    table3.style = "Table Grid"
    for i, (label, value) in enumerate(bank_fields):
        table3.rows[i].cells[0].text = label
        table3.rows[i].cells[1].text = str(value)

    doc.add_paragraph("")
    doc.add_paragraph("Declaration: I hereby declare that all information provided above is true.")
    doc.add_paragraph("\n\nSignature: ___________________          Date: ___________")
    doc.add_paragraph("(Applicant's Signature / Thumb Impression)")

    # Save
    filename = f"{form_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    filepath = Path(settings.GENERATED_DIR) / filename
    doc.save(str(filepath))
    return filename


def generate_form_pdf(form_type: str, data: dict) -> str:
    """Generate a filled PDF form and return filename."""
    filename = f"{form_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    filepath = Path(settings.GENERATED_DIR) / filename

    doc = SimpleDocTemplate(str(filepath), pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph(FORM_TEMPLATES.get(form_type, "Application Form"), styles["Title"]))
    elements.append(Paragraph("Ministry of Agriculture & Farmers Welfare", styles["Normal"]))
    elements.append(Spacer(1, 12))

    # Applicant table
    table_data = [["Field", "Value"]]
    fill_fields = ["name", "father_name", "dob", "id_number", "phone", "state", "district", "address",
                   "land_area", "crop_types", "bank_account", "ifsc_code"]
    labels = ["Full Name", "Father's Name", "Date of Birth", "Aadhaar/ID Number", "Mobile",
              "State", "District", "Address", "Land Area", "Crop Types", "Bank Account", "IFSC Code"]
    for label, key in zip(labels, fill_fields):
        table_data.append([label, str(data.get(key, ""))])

    t = Table(table_data, colWidths=[200, 300])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.black),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
    ]))
    elements.append(t)
    doc.build(elements)
    return filename
